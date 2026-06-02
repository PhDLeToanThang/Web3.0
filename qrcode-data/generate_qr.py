import os, sys, hashlib
from pathlib import Path
from openpyxl import load_workbook
import qrcode

BASE_DIR = Path("QR_Output")
QR1_DIR = BASE_DIR / "QR_FormID"
QR2_DIR = BASE_DIR / "QR_DataLink"
DATA_DIR = BASE_DIR / "Household_Data"

SERVER_PATH = "file:///SERVER/Survey_Data"
WEB_URL = "https://survey.xxx/data"

os.makedirs(QR1_DIR, exist_ok=True)
os.makedirs(QR2_DIR, exist_ok=True)
os.makedirs(DATA_DIR, exist_ok=True)

def read_excel(filepath):
    wb = load_workbook(filepath, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    return rows

def make_form_id(ma_phuong, ma_to, cccd, stt):
    raw = f"{ma_phuong}-{ma_to}-{cccd}-{stt:03d}"
    h = hashlib.md5(raw.encode()).hexdigest()[:6].upper()
    return f"{ma_phuong}-{ma_to}-{cccd}-{h}"

def make_folder_name(ma_phuong, ma_to, cccd, stt):
    return f"{ma_phuong}_{ma_to}_{cccd}_{stt:03d}"

def make_qr_image(data, filepath):
    qr = qrcode.QRCode(box_size=8, border=2)
    qr.add_data(data)
    qr.make(fit=True)
    img = qr.make_image(fill_color="black", back_color="white")
    img.save(filepath)

def make_qr2_path(ma_phuong, ma_to, cccd, stt):
    folder = make_folder_name(ma_phuong, ma_to, cccd, stt)
    local = str(DATA_DIR / folder)
    network = f"{SERVER_PATH}/{ma_phuong}_{ma_to}/{folder}"
    web = f"{WEB_URL}/{ma_phuong}_{ma_to}/{folder}"
    return {"local": local, "network": network, "web": web, "folder": folder}

stats = {"total": 0}

def process(rows):
    for idx, row in enumerate(rows, 1):
        if row[0] is None or str(row[0]).strip() == "":
            continue

        stt = int(row[0])
        ma_phuong = str(row[2]).strip() if row[2] else ""
        ma_to = str(row[4]).strip() if row[4] else ""
        cccd = str(row[5]).strip() if row[5] else ""
        ho_ten = str(row[6]).strip() if row[6] else ""

        if not ma_phuong or not ma_to or not cccd:
            print(f"[WARN] Dong {idx}: thieu du lieu, bo qua ({ho_ten})")
            continue

        form_id = make_form_id(ma_phuong, ma_to, cccd, stt)
        qr2 = make_qr2_path(ma_phuong, ma_to, cccd, stt)

        os.makedirs(qr2["local"], exist_ok=True)

        qr1_file = QR1_DIR / f"{form_id}.png"
        qr2_file = QR2_DIR / f"{form_id}_data.png"

        make_qr_image(form_id, qr1_file)
        make_qr_image(qr2["web"], qr2_file)

        stats["total"] += 1
        print(f"[OK] {stt:4d} | {ho_ten:20s} | FormID: {form_id}")

        crlf_file = Path(qr2["local"]) / "link_url.txt"
        crlf_file.write_text(
            f"URL truy cap: {qr2['web']}\n"
            f"Network path: {qr2['network']}\n"
            f"Folder local: {qr2['local']}\n",
            encoding="utf-8"
        )

        pdf_file = Path(qr2["local"]) / "dat_dummy_pdf_here.txt"
        pdf_file.write_text(
            "Dat file PDF survey, scan 3D, so do nha, so do,... vao thu muc nay.",
            encoding="utf-8"
        )

def report():
    print(f"\n{'='*60}")
    print("THONG KE:")
    print(f"  - Tong so ho da xu ly: {stats['total']}")
    print(f"  - QR FormID   : {QR1_DIR}/")
    print(f"  - QR DataLink : {QR2_DIR}/")
    print(f"  - Folder du lieu: {DATA_DIR}/")
    print(f"{'='*60}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Cach dung: python generate_qr.py <file_excel>")
        print("Vi du    : python generate_qr.py sample_data.xlsx")
        sys.exit(1)

    excel_file = sys.argv[1]
    if not os.path.exists(excel_file):
        print(f"[ERR] Khong tim thay file: {excel_file}")
        sys.exit(1)

    print(f"Dang doc file: {excel_file}")
    rows = read_excel(excel_file)
    print(f"Tim thay {len(rows)} dong du lieu.\n")

    process(rows)
    report()
