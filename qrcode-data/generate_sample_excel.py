from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = Workbook()
ws = wb.active
ws.title = "DanhSachHoGiaDinh"

headers = [
    "STT", "Phuong_Xa", "MaPhuong", "ToDP_Thon", "MaTo",
    "CCCD_ChuHo", "HoVaTen", "SoNha_DiaChi", "GhiChu"
]

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = thin_border

sample_data = [
    [1, "Phường An Khánh", "P01", "Tổ 5", "T05", "079201234567", "Nguyễn Văn A", "12/3 Đường A", ""],
    [2, "Phường An Khánh", "P01", "Tổ 5", "T05", "079201234568", "Trần Thị B", "45/7 Đường B", ""],
    [3, "Phường An Khánh", "P01", "Tổ 6", "T06", "079201234569", "Lê Văn C", "78 Đường C", ""],
    [4, "Phường Hòa Phú", "P02", "Tổ 2", "T02", "079201234570", "Phạm Thị D", "23/5 Đường D", "Có sổ đỏ"],
    [5, "Phường Hòa Phú", "P02", "Tổ 2", "T02", "079201234571", "Hoàng Văn E", "67/9 Đường E", ""],
    [6, "Phường Hòa Phú", "P02", "Tổ 3", "T03", "079201234572", "Đặng Thị F", "90 Đường F", ""],
    [7, "Phường Trường An", "P03", "Tổ 1", "T01", "079201234573", "Bùi Văn G", "34/12 Đường G", ""],
    [8, "Phường Trường An", "P03", "Tổ 1", "T01", "079201234574", "Đỗ Thị H", "56/78 Đường H", "Có scan 3D"],
    [9, "Phường Trường An", "P03", "Tổ 4", "T04", "079201234575", "Hồ Văn I", "12A Đường I", ""],
    [10, "Phường An Khánh", "P01", "Tổ 5", "T05", "079201234576", "Võ Thị K", "100/2 Đường K", ""],
]

for row_data in sample_data:
    ws.append(row_data)

ws.column_dimensions['A'].width = 6
ws.column_dimensions['B'].width = 18
ws.column_dimensions['C'].width = 10
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 8
ws.column_dimensions['F'].width = 16
ws.column_dimensions['G'].width = 20
ws.column_dimensions['H'].width = 22
ws.column_dimensions['I'].width = 15

wb.save("sample_data.xlsx")
print("Da tao sample_data.xlsx voi 10 ho gia dinh mau.")
